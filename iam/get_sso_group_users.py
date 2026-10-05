## This script requires credential to Management Account of the Org or a delegated admin for AWS IAM Identity Center
## This is to get all of Groups from AWS IAM Identity Center
## then get all the Users inside each Groups
## then get username of each user
import boto3
from botocore.config import Config
from botocore.exceptions import BotoCoreError, ClientError
import argparse
import pandas as pd
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

def main():
    p = argparse.ArgumentParser()
    p.add_argument("--profile", default="default")
    p.add_argument("--out", default="aws_sso_user_list.xlsx")
    args = p.parse_args()
    
    region = "us-east-1" ## this must be in us-east-1 since this is where we homed our Identity Center
    ##setup session in our profile and region
    try:
        session = boto3.Session(profile_name=args.profile, region_name=region)
    except BotoCoreError as err:
        raise SystemExit(f"Could not load AWS profile '{args.profile}': {err}")
    client_config = Config(retries={"mode": "adaptive", "max_attempts": 10})
    ##This client is needed to get list instances (list_instances)
    sso_admin_client = session.client('sso-admin', config=client_config)
    ##This client is needed to get groups (list_groups), get who is in groups (list_group_memberships)
    id_store_client = session.client('identitystore', config=client_config)
    # Get Identity Store ID, this will be used by Identity Store Client later
    try:
        instances = sso_admin_client.list_instances()["Instances"]
    except ClientError as err:
        raise SystemExit(f"Failed to get instance id with {err}")
    if not instances:
        raise SystemExit(f"No Identity Center instance found in {region} for profile {args.profile}")
    identity_store_id = instances[0]["IdentityStoreId"]
    ## prefetch all Users available, this makes it more efficient since boto3 is doing pagination for us
    user_dict = {}
    try:
        user_paginator = id_store_client.get_paginator('list_users')
        user_response_iterator = user_paginator.paginate(
            IdentityStoreId=identity_store_id
        )
        for user_response in user_response_iterator:
            for user in user_response["Users"]:
                user_dict[user["UserId"]] = user["UserName"]
    except (ClientError, BotoCoreError) as err:
        raise SystemExit(f"Failed to list users: {err}")            
    ##populate this to output later
    report_list = []
    group_paginator = id_store_client.get_paginator('list_groups')
    group_response_iterator = group_paginator.paginate(
        IdentityStoreId=identity_store_id,
    )
    group_member_paginator = id_store_client.get_paginator('list_group_memberships')
    try:
        for group_response in group_response_iterator:
            for group in group_response["Groups"]:
                group_id = group["GroupId"]
                display_name = group.get("DisplayName", "unknown")
                print(f"{group_id}:{display_name}")            
                group_info = {
                    "groupId": group_id,
                    "group": display_name,
                }
                member_response_iterator = group_member_paginator.paginate(
                    IdentityStoreId=identity_store_id,
                    GroupId=group_id,
                )
                member_count = 0
                try:
                    for member_response in member_response_iterator:
                        for member in member_response["GroupMemberships"]:
                            user_id = member['MemberId'].get('UserId')
                            if not user_id:
                                continue
                            member_count += 1
                            user_name = user_dict.get(user_id, "unknown")
                            report_list.append({**group_info, "user": user_name, "userId": user_id})
                except ClientError as err:
                    code = err.response["Error"]["Code"]
                    if code == "ResourceNotFoundException":
                        print(f"WARN: group {group_id} disappeared during run, skipping")
                        report_list.append({**group_info, "user": "group deleted during run", "userId": ""})
                        continue
                    raise SystemExit(f"Failed to list members of {display_name} ({group_id}): {err}")     
                if member_count == 0:                
                    report_list.append({**group_info, "user": "empty group", "userId": ""})
    except (ClientError, BotoCoreError) as err:
        raise SystemExit(f"Failed to list groups: {err}")
    columns=["groupId", "group", "user", "userId"]
    df = pd.DataFrame(report_list, columns=columns)
    # create Excel sheet and table
    try:
        with pd.ExcelWriter(args.out, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, header=True, sheet_name="AWS SSO Members")
            ws = writer.sheets["AWS SSO Members"]
            if len(df):
                ref = f"A1:{get_column_letter(len(columns))}{len(df) + 1}"
                table = Table(displayName="SSOMembers", ref=ref)
                table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
                ws.add_table(table)
    except PermissionError:
        raise SystemExit(f"Cannot write {args.out}: is it open in Excel? Close it and re-run.")
    except OSError as err:
        raise SystemExit(f"Failed to write {args.out}: {err}")
    print(f"Done. Wrote {len(df)} lines to {args.out}")

if __name__ == "__main__":
    main()
